"""
Skill: Work Paper Generator
Generates audit work papers from Excel or Word templates.
"""

import json
import pathlib


def _load_data(data_file: str) -> dict:
    with open(data_file, "r", encoding="utf-8") as fh:
        return json.load(fh)


def _fill_excel(template_path: str, data: dict, output_path: str) -> list[str]:
    import openpyxl

    wb = openpyxl.load_workbook(template_path)
    filled = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    for key, val in data.items():
                        placeholder = f"{{{{{key}}}}}"
                        if placeholder in cell.value:
                            cell.value = cell.value.replace(placeholder, str(val))
                            if key not in filled:
                                filled.append(key)
    wb.save(output_path)
    return filled


def _fill_word(template_path: str, data: dict, output_path: str) -> list[str]:
    from docx import Document

    doc = Document(template_path)
    filled = []
    for para in doc.paragraphs:
        for key, val in data.items():
            placeholder = f"{{{{{key}}}}}"
            if placeholder in para.text:
                for run in para.runs:
                    if placeholder in run.text:
                        run.text = run.text.replace(placeholder, str(val))
                        if key not in filled:
                            filled.append(key)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for key, val in data.items():
                        placeholder = f"{{{{{key}}}}}"
                        if placeholder in para.text:
                            for run in para.runs:
                                if placeholder in run.text:
                                    run.text = run.text.replace(placeholder, str(val))
                                    if key not in filled:
                                        filled.append(key)
    doc.save(output_path)
    return filled


def run(template_path: str, data_file: str, output_path: str) -> dict:
    try:
        data = _load_data(data_file)
    except Exception as exc:
        result = {"status": "error", "output_file": "", "fields_filled": [], "message": str(exc)}
        _write_json(output_path + ".result.json", result)
        return result

    template = pathlib.Path(template_path)
    suffix = template.suffix.lower()

    try:
        if suffix in {".xlsx", ".xls"}:
            filled = _fill_excel(template_path, data, output_path)
        elif suffix == ".docx":
            filled = _fill_word(template_path, data, output_path)
        else:
            result = {
                "status": "error",
                "output_file": "",
                "fields_filled": [],
                "message": f"Unsupported template type: {suffix}",
            }
            _write_json(output_path + ".result.json", result)
            return result
    except Exception as exc:
        result = {"status": "error", "output_file": "", "fields_filled": [], "message": str(exc)}
        _write_json(output_path + ".result.json", result)
        return result

    result = {"status": "ok", "output_file": str(output_path), "fields_filled": filled}
    _write_json(output_path + ".result.json", result)
    return result


def _write_json(json_path: str, data: dict) -> None:
    pathlib.Path(json_path).parent.mkdir(parents=True, exist_ok=True)
    with open(json_path, "w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False, indent=2)
